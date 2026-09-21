from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs'
SOURCE = ROOT / 'store'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    for border in st.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(28)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 2'].font.size = Pt(13)
code_style = doc.styles.add_style('CodeExample', 1)
code_style.font.name = 'Consolas'
code_style.font.size = Pt(8.5)
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing = 1.0

def p(text):
    doc.add_paragraph(text)

def h(text, level=1):
    doc.add_heading(text, level)

def code(text):
    for line in text.strip('\n').splitlines():
        doc.add_paragraph(line, 'CodeExample')
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def new(text):
    doc.add_page_break()
    h(text)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Cm(width)
    for i, text in enumerate(headers):
        t.rows[0].cells[i].text = text
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = 1
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + side)
                for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    el.set(qn('w:' + key), val)
                borders.append(el)
            pr.append(borders)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), 'E8EDF4' if ri == 0 else 'FFFFFF')
            pr.append(shading)
            margin = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), '100')
                el.set(qn('w:type'), 'dxa')
                margin.append(el)
            pr.append(margin)
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(3)
                for run in par.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

doc.add_paragraph('Tìm hiểu Redux Toolkit\ntrong React Native', 'Title')
doc.add_paragraph('Ứng dụng thực tế và demo lấy sản phẩm từ DummyJSON', 'Subtitle')
p('Báo cáo thực hành • Ngày 21 tháng 09 năm 2026')
p('Redux Toolkit giúp tổ chức state dùng chung và xử lý thay đổi dữ liệu theo một luồng rõ ràng. Báo cáo giới thiệu ứng dụng, cú pháp và cách tích hợp vào React Native, sau đó thực hành tải sản phẩm từ API và đánh dấu yêu thích. Mã nguồn demo nằm trong components/redux-toolkit.')
h('1 Redux Toolkit dùng để làm gì')
p('Redux lưu state trong store. Redux Toolkit là bộ công cụ chính thức giúp viết logic Redux với ít mã cấu hình hơn; React Redux kết nối store với component qua Provider và hooks. Cách tổ chức này dùng được trong React Native. [1]')
table(['Tình huống', 'Ví dụ state dùng chung'], [
    ('Giỏ hàng', 'Sản phẩm đã chọn, số lượng và tổng tiền hiển thị ở nhiều màn hình.'),
    ('Tài khoản', 'Thông tin hồ sơ và trạng thái đăng nhập mà các màn hình cùng đọc.'),
    ('Danh sách sản phẩm', 'Kết quả API, trạng thái đang tải, thông báo lỗi và yêu thích.'),
    ('Thiết lập ứng dụng', 'Ngôn ngữ hoặc bộ lọc được dùng ở nhiều component.'),
], [4.2, 12.8])
h('Khi nào nên sử dụng', 2)
p('Nên cân nhắc Redux Toolkit khi nhiều màn hình cần đọc hoặc thay đổi cùng dữ liệu, hoặc logic cập nhật khó theo dõi. Với ô nhập liệu hay trạng thái mở một hộp thoại chỉ dùng tại một component, useState thường đủ. Đây là cách lựa chọn phạm vi state cho ứng dụng, không phải yêu cầu đưa mọi dữ liệu vào Redux.')
h('Hai hướng xử lý API', 2)
p('Demo dùng createAsyncThunk để quan sát rõ từng trạng thái request. RTK Query là lựa chọn khác trong Redux Toolkit, hỗ trợ tải và lưu cache dữ liệu API, đồng thời sinh hooks cho React. Với ứng dụng có nhiều endpoint và nhu cầu cache, có thể dùng RTK Query để giảm phần logic tải dữ liệu phải tự viết. [2]')

new('2 Thành phần và cú pháp cơ bản')
table(['Thành phần', 'Vai trò trong demo'], [
    ('Store', 'Chứa state.products; được tạo bằng configureStore.'),
    ('Slice và reducer', 'Gom state ban đầu và logic cập nhật theo nhóm sản phẩm.'),
    ('Action và dispatch', 'Mô tả thay đổi; dispatch gửi action hoặc thunk đến store.'),
    ('Provider', 'Cấp store cho các component con.'),
    ('useSelector', 'Đọc dữ liệu và cập nhật component khi kết quả chọn thay đổi.'),
    ('useDispatch', 'Lấy hàm dispatch để gửi action từ sự kiện giao diện.'),
], [4.2, 12.8])
p('Luồng đồng bộ: người dùng nhấn nút → dispatch(action) → reducer tạo state tiếp theo → component đọc dữ liệu mới. createSlice tự sinh action creators từ reducers. Cú pháp cập nhật trực tiếp trong reducer của slice được Immer chuyển thành cập nhật bất biến. Không sửa state trực tiếp bên ngoài reducer. [3]')
h('Cài đặt và tạo store', 2)
code('npm install @reduxjs/toolkit react-redux')
p('Dự án sử dụng Expo 54.0.36, React Native 0.81.5, React 19.1.0, Redux Toolkit 2.12.0 và React Redux 9.3.0. Expo SDK 54 đi cùng React Native 0.81 và React 19.1.0; tài liệu đúng phiên bản đã được đối chiếu. [4]')
p('Tệp store/index.ts')
code((SOURCE / 'index.ts').read_text(encoding='utf-8'))
h('Hooks có kiểu TypeScript', 2)
code((SOURCE / 'hooks.ts').read_text(encoding='utf-8'))
p('RootState được suy luận từ store.getState, còn AppDispatch từ store.dispatch. Hooks được định kiểu giúp selector biết cấu trúc state và dispatch chấp nhận thunk. [5]')

new('3 Thiết kế demo sản phẩm')
p('Màn hình Home tự tải 10 sản phẩm khi mở. Mỗi dòng hiển thị ảnh, tên, giá và nút yêu thích. Component ProductSummary hiển thị tổng số sản phẩm đã tải và số sản phẩm yêu thích từ cùng store. Nút Tải lại sản phẩm gửi lại request.')
h('API sử dụng', 2)
code('GET https://dummyjson.com/products\n    ?limit=10&select=title,price,thumbnail')
p('URL thực tế nằm trên một dòng trong PRODUCTS_URL. limit giới hạn số bản ghi, select chọn trường cần lấy. API trả về object có mảng products và thông tin phân trang; demo chỉ sử dụng mảng products. Endpoint đọc sản phẩm này không yêu cầu đăng nhập. [6]')
h('Cấu trúc thư mục', 2)
code('components/redux-toolkit/\n  index.tsx\n  Header.tsx\n  ActionButton.tsx\n  ProductItem.tsx\n  ProductSummary.tsx\nstore/\n  index.ts\n  hooks.ts\n  productsSlice.ts\nscripts/verify-redux.cjs')
h('Kiểu dữ liệu và state ban đầu', 2)
source = (SOURCE / 'productsSlice.ts').read_text(encoding='utf-8')
code(source[source.index('export interface Product'):source.index('export const PRODUCTS_URL')])
p('items lưu dữ liệu từ server; favoriteIds lưu lựa chọn tại ứng dụng. status điều khiển trạng thái giao diện; error chứa thông báo cho người dùng. Không lưu Promise, Response hoặc Error trực tiếp vào Redux state.')
p('Yêu thích chỉ thay đổi trong bộ nhớ, không gọi API ghi dữ liệu. Tải lại sản phẩm giữ các ID đã thích; tải lại toàn bộ ứng dụng sẽ tạo state ban đầu. Demo giới hạn ở 10 sản phẩm và chưa triển khai phân trang hay lưu trữ lâu dài.')

new('4 Gọi API bằng createAsyncThunk')
p('Tệp productsSlice.ts khai báo fetchProducts như sau. Đoạn mã này lấy từ demo đang chạy, gồm kiểm tra HTTP, timeout 15 giây và chặn request trùng.')
code(source[source.index('export const PRODUCTS_URL'):source.index('const productsSlice')])
h('Ý nghĩa cú pháp', 2)
p('Product[] là kiểu kết quả; void cho biết thunk không nhận tham số; rejectValue: string định kiểu thông báo lỗi. Chuỗi products/fetchProducts là tiền tố action. Hàm bất đồng bộ trả về sản phẩm sẽ tạo fulfilled; rejectWithValue tạo rejected có payload lỗi do ứng dụng cung cấp. pending được phát khi bắt đầu thực thi. [7]')
p('fetch không tự coi mã HTTP lỗi là Promise bị reject, nên demo kiểm tra response.ok. finally luôn dọn timer. condition bỏ qua lần gọi mới khi status đang là loading; nút tải lại cũng bị vô hiệu ở giao diện. Các xử lý này giúp demo tải lại ổn định và không chờ vô hạn khi mạng chậm.')

new('5 Cập nhật state và kết nối React Native')
p('Cũng trong productsSlice.ts, reducers xử lý yêu thích đồng bộ; extraReducers nhận các action vòng đời do thunk phát. createAsyncThunk không tự tạo phần cập nhật state này. [7]')
code(source[source.index('const productsSlice'):])
h('Provider cấp store cho component', 2)
code((ROOT / 'app' / '_layout.tsx').read_text(encoding='utf-8'))
p('Tệp app/(tabs)/index.tsx render component demo:')
code((ROOT / 'app' / '(tabs)' / 'index.tsx').read_text(encoding='utf-8'))
p('Store được tạo một lần; Provider tại app/_layout.tsx cấp store cho tất cả màn hình.')

new('6 Sử dụng state trong giao diện')
p('ReduxToolkitDemo dùng useAppSelector để đọc state và useAppDispatch để gửi thunk hoặc action. Các trích đoạn dưới đây thể hiện phần kết nối; mã đầy đủ cùng StyleSheet có trong components/redux-toolkit/index.tsx.')
code("const dispatch = useAppDispatch();\nconst { items, status, error, favoriteIds } =\n  useAppSelector((state) => state.products);\n\nuseEffect(() => {\n  if (status === 'idle') void dispatch(fetchProducts());\n}, [dispatch, status]);")
p('Ví dụ gắn action vào sự kiện React Native:')
code("<Pressable\n  disabled={status === 'loading'}\n  onPress={() => void dispatch(fetchProducts())}>\n  <Text>Tải lại sản phẩm</Text>\n</Pressable>\n\n// Trong renderItem của FlatList\n<Pressable onPress={() => dispatch(toggleFavorite(item.id))}>\n  <Text>Yêu thích</Text>\n</Pressable>")
p('React Native dùng View, Text, Pressable, Image và FlatList. Sự kiện nhấn là onPress. Danh sách dùng keyExtractor với ID sản phẩm, không dùng vị trí phần tử làm khóa. Giá được hiển thị bằng item.price.toFixed(2).')
h('Component thứ hai đọc cùng store', 2)
code("const total = useAppSelector(\n  (state) => state.products.items.length,\n);\nconst favorites = useAppSelector(\n  (state) => state.products.favoriteIds.length,\n);\n\n<Text>Đã tải: {total} sản phẩm • Yêu thích: {favorites}</Text>")
p('Không cần truyền danh sách yêu thích từ ReduxToolkitDemo sang ProductSummary qua props. Khi toggleFavorite thay đổi favoriteIds, selector của ProductSummary nhận số lượng mới và component cập nhật. [3]')
h('Các trạng thái hiển thị', 2)
table(['State', 'Hành vi giao diện'], [
    ('idle', 'useEffect bắt đầu tải ở lần mở đầu.'),
    ('loading', 'Hiện ActivityIndicator và vô hiệu nút tải lại.'),
    ('succeeded', 'Hiện sản phẩm; nếu mảng rỗng, hiện Chưa có sản phẩm.'),
    ('failed', 'Hiện thông báo lỗi, cho tải lại; giữ danh sách cũ nếu có.'),
], [3.1, 13.9])

new('7 Chạy thử và kết quả kiểm tra')
h('Lệnh chạy tại thư mục gốc', 2)
code('npm install\nnpm run web\n# PowerShell chan npm.ps1: dung npm.cmd\n# Android khi da co thiet bi hoac emulator\nnpm run android')
p('Mở màn hình Home. Chờ danh sách tải xong, nhấn Yêu thích ở một sản phẩm và quan sát bộ đếm tăng; nhấn lại để bỏ thích. Nhấn Tải lại sản phẩm để thực hiện request mới. Cần Internet để gọi DummyJSON và tải ảnh.')
h('Kết quả đã thực hiện', 2)
table(['Kiểm tra', 'Kết quả'], [
    ('TypeScript và ESLint', 'Đạt, không có lỗi hoặc cảnh báo lint.'),
    ('Build web bằng Expo', 'Đạt; xuất bundle và 7 route vào dist.'),
    ('API thật qua thunk', 'Đạt; nhận 10 sản phẩm. Đầu tiên: Essence Mascara Lash Princess, giá 9.99.'),
    ('Vòng đời request', 'Đạt: pending, fulfilled và chặn request trùng.'),
    ('Yêu thích', 'Đạt: thêm, bỏ thích và state trước đó không bị sửa.'),
    ('Lỗi và tải lại', 'Đạt với mock HTTP 500, lỗi mạng, giữ dữ liệu cũ, retry và mảng rỗng.'),
], [5, 12])
p('Expo web đã khởi động tại http://localhost:8081. Các ca chức năng ở bảng là kiểm thử trực tiếp store/thunk bằng Node, không phải kiểm thử tự động bấm nút. Chưa kiểm tra trực tiếp giao diện Android/iOS hoặc tương tác trong trình duyệt.')
code('node scripts/verify-redux.cjs --live\nnpx tsc --noEmit --incremental false\nnpm run lint\nnpx expo export --platform web')
h('Tài liệu tham khảo', 2)
refs = [
    '[1] Redux Toolkit Getting Started — redux-toolkit.js.org/introduction/getting-started',
    '[2] RTK Query Overview — redux-toolkit.js.org/rtk-query/overview',
    '[3] Redux Toolkit Quick Start — redux-toolkit.js.org/tutorials/quick-start',
    '[4] Expo SDK 54 — docs.expo.dev/versions/v54.0.0/',
    '[5] TypeScript Quick Start — redux-toolkit.js.org/tutorials/typescript',
    '[6] DummyJSON Products — dummyjson.com/docs/products',
    '[7] createAsyncThunk — redux-toolkit.js.org/api/createAsyncThunk',
]
for ref in refs:
    par = doc.add_paragraph(ref)
    par.paragraph_format.space_after = Pt(3)
    for run in par.runs:
        run.font.size = Pt(9)

footer = sec.footer.paragraphs[0]
footer.alignment = 2
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Tìm hiểu Redux Toolkit trong React Native'
doc.core_properties.subject = 'Demo API DummyJSON và kết quả chạy thử'
doc.core_properties.author = ''
doc.save(OUT / 'Tim_hieu_Redux_Toolkit.docx')
print(OUT / 'Tim_hieu_Redux_Toolkit.docx')
